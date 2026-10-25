from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from csvftps import config as settings
from csvftps.config import Config, ConfigError, Rules
from csvftps.excel import ExcelError, cell_text, cell_value, read_rows
from csvftps.runner import RunError, run_once
from csvftps.transform import convert_file


def make_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    workbook.save(path)
    return path


SIZES = [
    ["Id", "Dealer", "City", "Group", "Size", "Price"],
    [1, "North", "Kyiv", "Phones", 1.5, 100],
    [2, "South", "Odesa", "Tablets", 7, 200],
    [3, "West", "Lviv", "Monitors", None, 300],
    [4, "East", "Kharkiv", "TV", "NULL", 900],
]


@pytest.mark.parametrize(
    ("value", "text"),
    [
        (None, ""),
        (12, "12"),
        (12.0, "12"),
        (7.5, "7.5"),
        (0.1, "0.1"),
        (True, "TRUE"),
        (datetime(2026, 10, 25), "2026-10-25"),
        (datetime(2026, 10, 25, 9, 30), "2026-10-25 09:30:00"),
        (date(2026, 1, 2), "2026-01-02"),
        ("text", "text"),
    ],
)
def test_cell_text(value: object, text: str) -> None:
    assert cell_text(value) == text


@pytest.mark.parametrize(
    ("text", "value"),
    [("", None), ("12", 12), ("-3", -3), ("7.5", 7.5), ("007", "007"), ("7,5", "7,5"), ("1e5", "1e5")],
)
def test_cell_value_keeps_text_that_is_not_plainly_a_number(text: str, value: object) -> None:
    assert cell_value(text) == value
    assert type(cell_value(text)) is type(value)


def test_xlsx_to_csv(tmp_path: Path) -> None:
    source = make_workbook(tmp_path / "in.xlsx", {"Data": SIZES})
    target = tmp_path / "out.csv"
    stats = convert_file(source, target, Rules(column="Size"))
    assert target.read_text(encoding="utf-8").splitlines() == [
        "Id;Dealer;City;Group;Size;Price",
        "1;North;Kyiv;Phones;None;100",
        "2;South;Odesa;Tablets;6-10 inch;200",
        "3;West;Lviv;Monitors;None;300",
        "4;East;Kharkiv;TV;None;900",
    ]
    assert (stats.converted, stats.empty) == (2, 2)


def test_csv_to_xlsx_keeps_numbers_and_bolds_headers(tmp_path: Path) -> None:
    source = tmp_path / "in.csv"
    source.write_text(
        "Id;Code;Group;Name Exchange;Size\n"
        "1;007;A;x;3\n"
        "2;=SUM(A1);B;y;12.5\n"
        "Id;Code;Group;Name Exchange;Size\n",
        encoding="utf-8",
    )
    target = tmp_path / "out.xlsx"
    convert_file(source, target, Rules())

    sheet = load_workbook(target).active
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [
        ["Id", "Code", "Group", "Name Exchange", "Size"],
        [1, "007", "A", "x", "2-5 inch"],
        [2, "=SUM(A1)", "B", "y", "> 10 inch"],
        ["Id", "Code", "Group", "Name Exchange", "Size"],
    ]
    assert sheet["B3"].data_type == "s"  # stored as text, never run as a formula
    assert sheet["A1"].font.bold and sheet["A4"].font.bold and not sheet["A2"].font.bold
    assert sheet.freeze_panes == "A2"


def test_xlsx_to_xlsx(tmp_path: Path) -> None:
    source = make_workbook(tmp_path / "in.xlsx", {"Data": SIZES})
    target = tmp_path / "out.xlsx"
    convert_file(source, target, Rules(column=4))
    rows = list(load_workbook(target).active.iter_rows(values_only=True))
    assert rows[1] == (1, "North", "Kyiv", "Phones", "None", 100)
    assert rows[2] == (2, "South", "Odesa", "Tablets", "6-10 inch", 200)


def test_sheet_by_name_and_number(tmp_path: Path) -> None:
    source = make_workbook(tmp_path / "in.xlsx", {"Notes": [["hello"]], "Data": SIZES})
    with read_rows(source) as rows:
        assert list(rows) == [["hello"]]
    with read_rows(source, "Data") as rows:
        assert next(rows)[0] == "Id"
    with read_rows(source, 2) as rows:
        assert next(rows)[0] == "Id"


@pytest.mark.parametrize(("sheet", "message"), [("Prices", "no sheet named 'Prices'"), (5, "no sheet 5")])
def test_missing_sheet(tmp_path: Path, sheet: int | str, message: str) -> None:
    source = make_workbook(tmp_path / "in.xlsx", {"Data": SIZES})
    with pytest.raises(ExcelError, match=message), read_rows(source, sheet) as rows:
        list(rows)


def test_trailing_blank_cells_and_rows_are_dropped(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path / "in.xlsx", {"S": [["a", "b", None], [None, None], ["c"], [None], [None, None]]}
    )
    with read_rows(source) as rows:
        assert list(rows) == [["a", "b"], [], ["c"]]


def test_not_a_workbook(tmp_path: Path) -> None:
    source = tmp_path / "fake.xlsx"
    source.write_text("Id;Size\n1;2\n", encoding="utf-8")
    config = Config(input=str(source), output=str(tmp_path / "out.csv"))
    with pytest.raises(RunError, match="fake.xlsx is not a valid Excel workbook"):
        run_once(config, do_upload=False)


def test_run_once_uploads_xlsx(project: Config, fake_ftps, monkeypatch) -> None:
    monkeypatch.setenv("CSVFTPS_PASSWORD", "secret")
    make_workbook(project.base_dir / "input.xlsx", {"Data": SIZES})
    project.input, project.output = "input.xlsx", "output.xlsx"
    project.rules.column = "Size"
    run_once(project)
    stored = fake_ftps.all_stored()
    assert list(stored) == ["prices.csv", "stock.csv", "output.xlsx"]
    assert stored["output.xlsx"][:2] == b"PK"  # a zip, i.e. a real workbook


@pytest.mark.parametrize(
    ("files", "message"),
    [
        ('input = "old.xls"', "old .xls workbooks are not supported"),
        ('output = "out.xlsm"', "cannot be written; use .xlsx"),
    ],
)
def test_unsupported_excel_files(tmp_path: Path, files: str, message: str) -> None:
    path = tmp_path / "csvftps.toml"
    path.write_text(f"[files]\n{files}\n", encoding="utf-8")
    with pytest.raises(ConfigError, match=message):
        settings.load(path)


def test_sheet_setting_round_trip(tmp_path: Path) -> None:
    for sheet in ("Prices", 2, ""):
        config = Config()
        config.rules.sheet = sheet
        assert settings.load(settings.save(config, tmp_path / "c.toml")).rules.sheet == sheet
    path = tmp_path / "bad.toml"
    path.write_text("[excel]\nsheet = 0\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="number starting at 1"):
        settings.load(path)


@pytest.mark.skipif(sys.platform != "win32", reason="only Windows refuses to replace a file that is open")
def test_output_open_in_another_program(project: Config) -> None:
    project.output = "output.xlsx"
    project.output_path.write_bytes(b"old")
    with project.output_path.open("rb"), pytest.raises(RunError, match="cannot access output.xlsx"):
        run_once(project, do_upload=False)
    assert project.output_path.read_bytes() == b"old"
    assert not list(project.base_dir.glob(".*.tmp"))
